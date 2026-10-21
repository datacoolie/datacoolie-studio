from __future__ import annotations

import json
from io import BytesIO
from types import SimpleNamespace

import pytest


def _source(uri: str = "metadata.json") -> SimpleNamespace:
    return SimpleNamespace(
        id=7,
        environment_id=3,
        source_kind="metadata",
        uri=uri,
        label="metadata",
        environment=None,
        source_config_json="{}",
        storage_provider="local",
        read_check_result_json=None,
    )


def _metadata_payload() -> dict:
    return {
        "$schema": "https://example.test/datacoolie-metadata.json",
        "connections": [
            {
                "name": "lake",
                "connection_id": "connection-explicit",
                "connection_type": "file",
                "format": "parquet",
            }
        ],
        "dataflows": [
            {
                "name": "orders",
                "dataflow_id": "dataflow-explicit",
                "workspace_id": "workspace-1",
                "group_number": 2,
                "execution_order": 4,
                "source": {"connection_name": "lake", "table": "raw.orders"},
                "destination": {"connection_name": "lake", "table": "silver.orders"},
            }
        ],
        "schema_hints": [
            {
                "connection_id": "connection-explicit",
                "table_name": "raw.orders",
                "hints": [{"column_name": "order_id", "data_type": "integer"}],
            }
        ],
    }


def test_strict_reader_unwraps_metadata_and_preserves_top_level_extras():
    from datacoolie_studio.domains.metadata.reader import (
        read_metadata_document_bytes,
    )

    content = json.dumps({"metadata": _metadata_payload()}).encode("utf-8")
    document = read_metadata_document_bytes("metadata.json", content, strict=True)

    assert document.wrapper == "metadata"
    assert document.payload["connections"][0]["name"] == "lake"
    assert document.extras == {"$schema": "https://example.test/datacoolie-metadata.json"}

    opaque = read_metadata_document_bytes(
        "metadata.json", json.dumps({"connections": {}}).encode("utf-8"), strict=True,
    )
    assert opaque.payload == {"connections": {}}


def test_strict_reader_rejects_excel_without_metadata_sheets():
    import openpyxl

    from datacoolie_studio.domains.metadata.reader import (
        MetadataReadError,
        read_metadata_document_bytes,
    )

    workbook = openpyxl.Workbook()
    workbook.active.title = "notes"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()

    with pytest.raises(MetadataReadError, match="none of the supported sheets"):
        read_metadata_document_bytes("metadata.xlsx", stream.getvalue(), strict=True)


def test_editor_round_trip_keeps_wrapper_extras_and_authored_runtime_ids():
    from datacoolie_studio.domains.metadata.editor import load_editor_document_from_raw
    from datacoolie_studio.domains.metadata.document_normalization import (
        editor_document_to_raw_metadata,
    )
    from datacoolie_studio.domains.metadata.reader import read_metadata_document_bytes

    source = _source()
    raw = {"metadata": _metadata_payload()}
    document_info = read_metadata_document_bytes(
        source.uri,
        json.dumps(raw).encode("utf-8"),
        strict=True,
    )
    editor_document = load_editor_document_from_raw(
        source,
        document_info.payload,
        {"content_hash": "source-hash"},
        document_info=document_info,
    )

    saved = editor_document_to_raw_metadata(editor_document)

    assert saved == raw
    assert editor_document["metadata_document"] == {
        "wrapper": "metadata",
        "extras": {"$schema": "https://example.test/datacoolie-metadata.json"},
    }


def test_modular_fragment_round_trip_preserves_list_and_record_shapes():
    from datacoolie_studio.domains.metadata.document_normalization import (
        editor_document_to_source_metadata,
    )
    from datacoolie_studio.domains.metadata.document_serialization import (
        serialize_editor_document_bytes,
    )
    from datacoolie_studio.domains.metadata.editor import load_editor_document_from_raw
    from datacoolie_studio.domains.metadata.reader import (
        MetadataReadError,
        read_metadata_document_bytes,
    )

    list_source = _source("metadata/dataflows.json")
    list_content = b'[{"name":"orders"}]'
    list_info = read_metadata_document_bytes(list_source.uri, list_content)
    assert list_info.fragment_section == "dataflows"
    assert list_info.fragment_shape == "list"

    list_document = load_editor_document_from_raw(
        list_source,
        list_info.payload,
        {"content_hash": "list-hash"},
        document_info=list_info,
    )
    list_document["sheets"]["dataflows"]["rows"][0]["name"] = "renamed"
    assert json.loads(
        serialize_editor_document_bytes(
            list_document,
            list_source.uri,
            original_bytes=list_content,
        )
    ) == [{"name": "renamed"}]

    record_source = _source("metadata/dataflows.json")
    record_content = b'{"name":"orders"}'
    record_info = read_metadata_document_bytes(record_source.uri, record_content)
    assert record_info.fragment_shape == "record"
    record_document = load_editor_document_from_raw(
        record_source,
        record_info.payload,
        {"content_hash": "record-hash"},
        document_info=record_info,
    )
    assert editor_document_to_source_metadata(record_document) == {"name": "orders"}
    record_document["sheets"]["dataflows"]["rows"].append({"name": "second"})
    with pytest.raises(MetadataReadError, match="single-record metadata fragment"):
        serialize_editor_document_bytes(
            record_document,
            record_source.uri,
            original_bytes=record_content,
        )


def test_yaml_round_trip_preserves_comments_and_unknown_top_level_fields():
    import yaml

    from datacoolie_studio.domains.metadata.document_serialization import (
        serialize_editor_document_bytes,
    )
    from datacoolie_studio.domains.metadata.editor import load_editor_document_from_raw
    from datacoolie_studio.domains.metadata.reader import read_metadata_document_bytes

    source = _source("metadata.yaml")
    content = (
        b"# Keep this file comment\n"
        b"connections:\n"
        b"  - name: lake  # Keep the row comment\n"
        b"dataflows: []\n"
        b"schema_hints: []\n"
        b"custom:\n"
        b"  owner: analytics\n"
    )
    info = read_metadata_document_bytes(source.uri, content)
    document = load_editor_document_from_raw(
        source,
        info.payload,
        {"content_hash": "yaml-hash"},
        document_info=info,
    )
    document["sheets"]["connections"]["rows"][0]["name"] = "warehouse"

    serialized = serialize_editor_document_bytes(
        document,
        source.uri,
        original_bytes=content,
    )
    assert b"# Keep this file comment" in serialized
    assert b"# Keep the row comment" in serialized
    assert yaml.safe_load(serialized)["connections"][0]["name"] == "warehouse"
    assert yaml.safe_load(serialized)["custom"] == {"owner": "analytics"}


def test_xlsx_round_trip_preserves_unknown_sheets_styles_and_comments():
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill

    from datacoolie_studio.domains.metadata.document_serialization import (
        serialize_editor_document_bytes,
    )
    from datacoolie_studio.domains.metadata.editor import load_editor_document_from_raw
    from datacoolie_studio.domains.metadata.reader import read_metadata_document_bytes

    workbook = openpyxl.Workbook()
    connections = workbook.active
    connections.title = "connections"
    connections.append(["name"])
    connections.append(["lake"])
    connections["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    connections["A2"].comment = Comment("Keep this comment", "DataCoolie")
    notes = workbook.create_sheet("notes")
    notes["A1"] = "do not remove"
    notes.sheet_state = "hidden"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    original = stream.getvalue()

    source = _source("metadata.xlsx")
    info = read_metadata_document_bytes(source.uri, original)
    document = load_editor_document_from_raw(
        source,
        info.payload,
        {"content_hash": "xlsx-hash"},
        document_info=info,
    )
    document["sheets"]["connections"]["rows"][0]["name"] = "warehouse"
    serialized = serialize_editor_document_bytes(
        document,
        source.uri,
        original_bytes=original,
    )

    updated = openpyxl.load_workbook(BytesIO(serialized), data_only=False)
    assert "notes" in updated.sheetnames
    assert updated["notes"]["A1"].value == "do not remove"
    assert updated["notes"].sheet_state == "hidden"
    assert updated["connections"]["A2"].value == "warehouse"
    assert updated["connections"]["A1"].fill.fgColor.rgb == "00FFFF00"
    assert updated["connections"]["A2"].comment.text == "Keep this comment"
    updated.close()


def test_schema_hints_accept_connection_id_without_connection_name():
    from datacoolie_studio.domains.metadata.document_normalization import (
        editor_document_to_raw_metadata,
    )

    document = {
        "sheets": {
            "connections": {
                "rows": [{"name": "lake", "connection_id": "connection-explicit"}]
            },
            "dataflows": {"rows": []},
            "schema_hints": {
                "rows": [
                    {
                        "connection_id": "connection-explicit",
                        "table_name": "orders",
                        "column_name": "id",
                    }
                ]
            },
        }
    }

    saved = editor_document_to_raw_metadata(document)
    assert saved["schema_hints"] == [
        {
            "connection_id": "connection-explicit",
            "table_name": "orders",
            "hints": [{"column_name": "id"}],
        }
    ]
