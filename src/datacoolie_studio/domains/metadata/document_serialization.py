"""Pure metadata document serializers for JSON, YAML, and XLSX outputs."""

from __future__ import annotations

import json
import tempfile
from collections.abc import Mapping, MutableMapping, MutableSequence
from copy import deepcopy
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

import yaml

from datacoolie_studio.domains.metadata.document_normalization import (
    editor_document_to_source_metadata,
    is_default_omitted_value,
)
from datacoolie_studio.domains.metadata.document_primitives import (
    is_empty,
    parse_cell_value,
    sheet_rows,
)
from datacoolie_studio.domains.metadata.document_schema import (
    FILE_RUNTIME_FIELDS,
    REQUIRED_SHEETS,
)
from datacoolie_studio.domains.metadata.reader import MetadataReadError


def serialize_editor_document(
    document: dict[str, Any],
    path: Path,
    *,
    original_bytes: bytes | None = None,
) -> str:
    raw = editor_document_to_source_metadata(document)
    suffix = path.suffix.lower()
    if suffix in {".yaml", ".yml"}:
        return _serialize_yaml(raw, original_bytes=original_bytes)
    return json.dumps(raw, indent=2, ensure_ascii=False)


def serialize_editor_document_bytes(
    document: dict[str, Any],
    uri: str,
    *,
    original_bytes: bytes | None = None,
) -> bytes:
    suffix = Path(uri).suffix.lower()
    if suffix not in {".xlsx", ".xls"}:
        return serialize_editor_document(
            document,
            Path(uri),
            original_bytes=original_bytes,
        ).encode("utf-8")
    if suffix == ".xls":
        raise MetadataReadError(
            "Unsupported metadata extension '.xls'; DataCoolie supports '.xlsx'"
        )
    with tempfile.TemporaryDirectory(prefix="datacoolie-metadata-") as directory:
        target = Path(directory) / f"metadata{suffix}"
        write_xlsx_document(document, target, original_bytes=original_bytes)
        return target.read_bytes()


def write_xlsx_document(
    document: dict[str, Any],
    path: Path,
    *,
    original_bytes: bytes | None = None,
) -> None:
    import openpyxl

    if original_bytes is None:
        workbook = openpyxl.Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)
    else:
        try:
            workbook = openpyxl.load_workbook(
                BytesIO(original_bytes),
                data_only=False,
                keep_links=True,
            )
        except Exception as exc:
            raise MetadataReadError(
                "Cannot preserve the existing XLSX workbook while saving metadata"
            ) from exc
        _assert_xlsx_round_trip_safe(workbook)

    sheets = document.get("sheets") or {}
    for sheet_name in REQUIRED_SHEETS:
        sheet_data = sheets.get(sheet_name) or {}
        excluded_fields = FILE_RUNTIME_FIELDS[sheet_name]
        rows = sheet_rows(sheets, sheet_name)
        columns = [
            column.get("key")
            for column in sheet_data.get("columns", [])
            if (
                column.get("key")
                and not str(column.get("key")).startswith("__")
                and (
                    column.get("key") not in excluded_fields
                    or any(
                        row.get(f"__studio_authored_{column.get('key')}") is True
                        for row in rows
                    )
                )
                and xlsx_column_has_saved_values(rows, str(column.get("key")))
            )
        ]
        if sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            _clear_xlsx_values(worksheet)
        else:
            worksheet = workbook.create_sheet(sheet_name)
        _write_xlsx_rows(worksheet, columns, rows)
    workbook.save(path)


def _serialize_yaml(raw: Any, *, original_bytes: bytes | None) -> str:
    try:
        yaml_writer = _round_trip_yaml()
    except ImportError:
        if original_bytes is not None:
            raise MetadataReadError(
                "ruamel.yaml is required to preserve comments in an existing YAML source"
            )
        return yaml.safe_dump(raw, sort_keys=False, allow_unicode=True)

    if original_bytes is not None:
        try:
            original = yaml_writer.load(original_bytes.decode("utf-8"))
        except Exception as exc:
            raise MetadataReadError(
                "Cannot preserve the existing YAML document while saving metadata"
            ) from exc
        raw = _merge_yaml_round_trip(original, raw)

    yaml_writer.default_flow_style = False
    buffer = StringIO()
    yaml_writer.dump(raw, buffer)
    return buffer.getvalue()


def _round_trip_yaml() -> Any:
    from ruamel.yaml import YAML

    yaml_writer = YAML(typ="rt")
    yaml_writer.preserve_quotes = True
    return yaml_writer


def _merge_yaml_round_trip(original: Any, desired: Any) -> Any:
    """Patch managed values into a round-trip YAML tree.

    The editor owns the canonical metadata sections, while comments and
    unrelated keys belong to the source document.  Updating existing mapping
    keys in place keeps their comments and formatting metadata.  Sequences are
    rebuilt by position so comments on unchanged row keys survive ordinary
    edits without pretending that row identity is globally available here.
    """

    if isinstance(original, MutableMapping) and isinstance(desired, Mapping):
        for key, value in desired.items():
            if key in original:
                original[key] = _merge_yaml_round_trip(original[key], value)
            else:
                original[key] = deepcopy(value)
        return original
    if isinstance(original, MutableSequence) and isinstance(desired, list):
        existing = list(original)
        merged = [
            _merge_yaml_round_trip(existing[index], value)
            if index < len(existing)
            else deepcopy(value)
            for index, value in enumerate(desired)
        ]
        original[:] = merged
        return original
    return deepcopy(desired)


def _assert_xlsx_round_trip_safe(workbook: Any) -> None:
    supported = set(REQUIRED_SHEETS)
    if not supported.intersection(workbook.sheetnames):
        raise MetadataReadError(
            "Excel metadata file contains none of the supported sheets"
        )
    if getattr(workbook, "vba_archive", None) is not None:
        raise MetadataReadError(
            "XLSX metadata with macros cannot be edited by DataCoolie"
        )
    if getattr(workbook, "_external_links", None):
        raise MetadataReadError(
            "XLSX metadata with external links cannot be edited safely"
        )
    for sheet_name in supported.intersection(workbook.sheetnames):
        worksheet = workbook[sheet_name]
        if worksheet.merged_cells.ranges:
            raise MetadataReadError(
                f"XLSX metadata sheet '{sheet_name}' contains merged cells"
            )
        if getattr(worksheet, "_charts", None) or getattr(worksheet, "_images", None):
            raise MetadataReadError(
                f"XLSX metadata sheet '{sheet_name}' contains drawings"
            )
        if getattr(worksheet, "_pivots", None) or getattr(worksheet, "tables", None):
            raise MetadataReadError(
                f"XLSX metadata sheet '{sheet_name}' contains unsupported table features"
            )
        for row in worksheet.iter_rows():
            if any(cell.data_type == "f" for cell in row):
                raise MetadataReadError(
                    f"XLSX metadata sheet '{sheet_name}' contains formulas"
                )


def _clear_xlsx_values(worksheet: Any) -> None:
    for row in worksheet.iter_rows():
        for cell in row:
            cell.value = None


def _write_xlsx_rows(
    worksheet: Any,
    columns: list[str],
    rows: list[dict[str, Any]],
) -> None:
    for column_index, column in enumerate(columns, start=1):
        worksheet.cell(row=1, column=column_index).value = column
    for row_index, row in enumerate(rows, start=2):
        for column_index, column in enumerate(columns, start=1):
            worksheet.cell(row=row_index, column=column_index).value = xlsx_cell_value(
                row.get(column),
                str(column),
            )


def xlsx_cell_value(value: Any, key: str | None = None) -> Any:
    if key is not None and is_default_omitted_value(key, value):
        return None
    parsed = parse_cell_value(value)
    if isinstance(parsed, (dict, list)):
        return json.dumps(parsed, ensure_ascii=False)
    return parsed


def xlsx_column_has_saved_values(rows: list[dict[str, Any]], key: str) -> bool:
    return any(
        not is_empty(row.get(key))
        and not is_default_omitted_value(key, row.get(key))
        for row in rows
    )
